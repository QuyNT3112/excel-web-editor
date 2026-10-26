import os
import io
from flask import Flask, render_template, request, jsonify, send_file
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
import webbrowser
import threading
import sys

if getattr(sys, 'frozen', False):
    bundle_dir = getattr(sys, '_MEIPASS', os.path.dirname(os.path.abspath(__file__)))
    app = Flask(
        __name__,
        template_folder=os.path.join(bundle_dir, 'templates'),
        static_folder=os.path.join(bundle_dir, 'static')
    )
    base_dir = os.path.dirname(sys.executable)
else:
    app = Flask(__name__)
    base_dir = os.path.dirname(os.path.abspath(__file__))

app.config['MAX_CONTENT_LENGTH'] = 32 * 1024 * 1024  # 32MB file limit

@app.route('/')
def index():
    initial_path = app.config.get('INITIAL_FILE_PATH')
    initial_filename = None
    initial_data_json = None
    if initial_path and os.path.isfile(initial_path):
        initial_filename = os.path.basename(initial_path)
        try:
            import json
            data = parse_spreadsheet_data(initial_path, initial_filename, filepath=initial_path)
            initial_data_json = json.dumps(data).replace('</script>', '<\\/script>').replace('<!--', '<\\!--')
        except Exception as e:
            print(f"[Warn] Pre-parsing initial file failed: {e}")
            
    return render_template('index.html', initial_filename=initial_filename, initial_data_json=initial_data_json)

@app.route('/health')
def health():
    return jsonify({'status': 'ok', 'app': 'Excel Web Pro', 'version': '2.1.0'})

def get_target_file_from_args():
    if len(sys.argv) <= 1:
        return None
    for arg in sys.argv[1:]:
        if not arg or arg.startswith('-'):
            continue
        cleaned = arg.strip('\'"')
        if os.path.isfile(cleaned):
            return os.path.abspath(cleaned)
    combined = " ".join([a for a in sys.argv[1:] if not a.startswith('-')]).strip('\'"')
    if os.path.isfile(combined):
        return os.path.abspath(combined)
    return None

cli_initial = get_target_file_from_args()
if cli_initial:
    app.config['INITIAL_FILE_PATH'] = cli_initial

def parse_spreadsheet_data(file_source, filename, filepath=None):
    """
    Parses a spreadsheet with high performance using iter_rows.
    Supports CSV and Excel (.xlsx, .xlsm, .xltx, etc.).
    Returns dict with filename, filepath, sheets, styles, colWidths, sheetNames.
    """
    ext = os.path.splitext(filename)[1].lower()
    
    if ext == '.csv':
        import csv
        if isinstance(file_source, str) and os.path.isfile(file_source):
            with open(file_source, 'r', encoding='utf-8-sig', errors='ignore') as f:
                content = f.read()
        elif hasattr(file_source, 'read'):
            raw = file_source.read()
            content = raw.decode('utf-8-sig', errors='ignore') if isinstance(raw, bytes) else raw
        elif isinstance(file_source, bytes):
            content = file_source.decode('utf-8-sig', errors='ignore')
        else:
            content = str(file_source)

        reader = csv.reader(io.StringIO(content))
        sheet_data = [[c for c in row] for row in reader]
        
        row_count = max(len(sheet_data), 100)
        col_count = max(max((len(r) for r in sheet_data), default=0), 26)
        
        padded_data = []
        for r in range(row_count):
            if r < len(sheet_data):
                row = sheet_data[r]
                padded_data.append(row + ["" for _ in range(col_count - len(row))])
            else:
                padded_data.append(["" for _ in range(col_count)])
                
        return {
            'filename': filename,
            'filepath': filepath or (file_source if isinstance(file_source, str) else None),
            'sheets': {'Sheet1': padded_data},
            'styles': {},
            'colWidths': {},
            'merges': {},
            'sheetNames': ['Sheet1']
        }

    # Excel file loading (.xlsx, .xlsm, etc.)
    if isinstance(file_source, str) and os.path.isfile(file_source):
        with open(file_source, 'rb') as f:
            excel_bytes = f.read()
        wb_source = io.BytesIO(excel_bytes)
        file_path_recorded = file_source
    elif hasattr(file_source, 'read'):
        raw = file_source.read()
        wb_source = io.BytesIO(raw) if isinstance(raw, bytes) else file_source
        file_path_recorded = filepath
    elif isinstance(file_source, bytes):
        wb_source = io.BytesIO(file_source)
        file_path_recorded = filepath
    else:
        wb_source = file_source
        file_path_recorded = filepath

    try:
        wb = openpyxl.load_workbook(wb_source, data_only=False)
    except Exception:
        if hasattr(wb_source, 'seek'):
            wb_source.seek(0)
        wb = openpyxl.load_workbook(wb_source, data_only=True)
    
    sheets = {}
    styles = {}
    col_widths = {}
    images = {}
    merges = {}
    import base64
    import time
    
    for sheet_name in wb.sheetnames:
        ws = wb[sheet_name]
        actual_rows = ws.max_row or 0
        actual_cols = ws.max_column or 0
        target_rows = max(actual_rows, 100)
        target_cols = max(actual_cols, 26)
        
        data_matrix = []
        sheet_styles = {}
        
        # Extract merged cells from worksheet
        sheet_merges = []
        if hasattr(ws, 'merged_cells') and ws.merged_cells:
            for m in list(ws.merged_cells.ranges):
                sheet_merges.append(str(m))
        merges[sheet_name] = sheet_merges
        
        if actual_rows > 0 and actual_cols > 0:
            for r_idx, row in enumerate(ws.iter_rows(min_row=1, max_row=actual_rows, min_col=1, max_col=actual_cols)):
                row_vals = []
                for c_idx, cell in enumerate(row):
                    val = cell.value
                    if val is None:
                        row_vals.append("")
                    elif isinstance(val, (int, float)):
                        row_vals.append(val)
                    else:
                        row_vals.append(str(val))
                    
                    # Extract styles only if styled
                    if cell.has_style:
                        c_style = {}
                        if cell.font:
                            if cell.font.bold: c_style['bold'] = True
                            if cell.font.italic: c_style['italic'] = True
                            if cell.font.underline: c_style['underline'] = True
                            if cell.font.color and cell.font.color.rgb and isinstance(cell.font.color.rgb, str):
                                rgb = cell.font.color.rgb
                                c_style['color'] = f"#{rgb[2:]}" if len(rgb) == 8 else f"#{rgb}"
                        if cell.fill and cell.fill.fill_type == 'solid' and cell.fill.start_color:
                            rgb = cell.fill.start_color.rgb
                            if rgb and isinstance(rgb, str):
                                c_style['bg'] = f"#{rgb[2:]}" if len(rgb) == 8 else f"#{rgb}"
                        if cell.alignment:
                            if cell.alignment.horizontal:
                                c_style['align'] = cell.alignment.horizontal
                            if cell.alignment.vertical:
                                c_style['valign'] = cell.alignment.vertical
                        if c_style:
                            sheet_styles[f"{r_idx},{c_idx}"] = c_style
                            
                if len(row_vals) < target_cols:
                    row_vals.extend([""] * (target_cols - len(row_vals)))
                data_matrix.append(row_vals)
                
        # Fast padding of remaining rows
        empty_row = [""] * target_cols
        while len(data_matrix) < target_rows:
            data_matrix.append(list(empty_row))
            
        sheets[sheet_name] = data_matrix
        styles[sheet_name] = sheet_styles
        
        ws_widths = {}
        for col_letter, col_dim in ws.column_dimensions.items():
            if col_dim.width:
                ws_widths[col_letter] = col_dim.width
        col_widths[sheet_name] = ws_widths

        # Extract images from worksheet
        sheet_images = []
        if hasattr(ws, '_images') and ws._images:
            for idx, img in enumerate(ws._images):
                try:
                    raw_bytes = img._data()
                    b64 = base64.b64encode(raw_bytes).decode('ascii')
                    fmt = getattr(img, 'format', 'png') or 'png'
                    data_url = f"data:image/{fmt.lower()};base64,{b64}"

                    col_idx = 1
                    row_idx = 1
                    if hasattr(img, 'anchor') and hasattr(img.anchor, '_from'):
                        col_idx = getattr(img.anchor._from, 'col', 0)
                        row_idx = getattr(img.anchor._from, 'row', 0)
                    elif isinstance(getattr(img, 'anchor', None), str):
                        from openpyxl.utils import coordinate_to_tuple
                        cr, cc = coordinate_to_tuple(img.anchor)
                        row_idx = cr - 1
                        col_idx = cc - 1

                    w = getattr(img, 'width', 220) or 220
                    h = getattr(img, 'height', 160) or 160

                    sheet_images.append({
                        'id': f"img-{idx}-{int(time.time() * 1000)}",
                        'src': data_url,
                        'name': f"Image_{idx + 1}.{fmt.lower()}",
                        'col': col_idx,
                        'row': row_idx,
                        'x': col_idx * 110 + 40,
                        'y': row_idx * 26 + 30,
                        'width': int(w),
                        'height': int(h),
                        'zIndex': 10 + idx
                    })
                except Exception as img_err:
                    print(f"[Warn] Could not extract image {idx}: {img_err}")
        images[sheet_name] = sheet_images
        
    return {
        'filename': filename,
        'filepath': file_path_recorded,
        'sheets': sheets,
        'styles': styles,
        'colWidths': col_widths,
        'merges': merges,
        'images': images,
        'sheetNames': wb.sheetnames
    }

@app.route('/upload', methods=['POST'])
def upload_file():
    if 'file' not in request.files:
        return jsonify({'error': 'No file part in the request'}), 400
    
    file = request.files['file']
    if file.filename == '':
        return jsonify({'error': 'No file selected'}), 400
    
    allowed_exts = ('.xlsx', '.xlsm', '.xltx', '.xltm', '.xls', '.csv')
    if not file.filename.lower().endswith(allowed_exts):
        return jsonify({'error': 'Unsupported file format. Please upload an Excel (.xlsx, .xls) or CSV file.'}), 400
    
    try:
        data = parse_spreadsheet_data(file, file.filename)
        return jsonify(data)
    except Exception as e:
        return jsonify({'error': f'Failed to process file: {str(e)}'}), 500

@app.route('/api/initial-file', methods=['GET'])
def get_initial_file():
    initial_path = app.config.get('INITIAL_FILE_PATH')
    if not initial_path or not os.path.isfile(initial_path):
        return jsonify({'has_file': False})
    
    return jsonify({
        'has_file': True,
        'filename': os.path.basename(initial_path),
        'filepath': initial_path
    })

@app.route('/api/load-initial-file', methods=['GET'])
def load_initial_file():
    initial_path = app.config.get('INITIAL_FILE_PATH')
    if not initial_path or not os.path.isfile(initial_path):
        return jsonify({'error': 'No initial file specified or file not found'}), 404
    
    filename = os.path.basename(initial_path)
    try:
        data = parse_spreadsheet_data(initial_path, filename, filepath=initial_path)
        return jsonify(data)
    except Exception as e:
        # Fallback to serving raw bytes for client-side SheetJS parsing
        return jsonify({
            'fallback_to_raw': True,
            'raw_url': '/api/initial-file/raw',
            'filename': filename,
            'filepath': initial_path,
            'error': str(e)
        })

@app.route('/api/initial-file/raw', methods=['GET'])
def get_initial_file_raw():
    initial_path = app.config.get('INITIAL_FILE_PATH')
    if not initial_path or not os.path.isfile(initial_path):
        return "File not found", 404
    return send_file(initial_path, as_attachment=False)

@app.route('/api/save_in_place', methods=['POST'])
def save_in_place():
    data = None
    if request.is_json:
        data = request.get_json()
    else:
        import json
        raw = request.form.get('payload') or request.form.get('data')
        if raw:
            try:
                data = json.loads(raw)
            except Exception:
                data = None

    if not data or 'sheets' not in data:
        return jsonify({'error': 'Invalid request data'}), 400

    filepath = data.get('filepath') or app.config.get('INITIAL_FILE_PATH')
    if not filepath:
        return jsonify({'error': 'No filepath specified'}), 400

    parent_dir = os.path.dirname(filepath)
    if parent_dir and not os.path.exists(parent_dir):
        os.makedirs(parent_dir, exist_ok=True)

    try:
        wb, _ = generate_openpyxl_workbook(data)
        wb.save(filepath)
        return jsonify({
            'success': True,
            'saved_to': filepath,
            'filename': os.path.basename(filepath)
        })
    except Exception as e:
        return jsonify({'error': f'Failed to save file: {str(e)}'}), 500

EXPORT_DIR = os.path.join(base_dir, 'exports')
os.makedirs(EXPORT_DIR, exist_ok=True)

def generate_openpyxl_workbook(data):
    filename = data.get('filename', 'export.xlsx')
    import re, unicodedata
    name_str = str(filename).replace('đ', 'd').replace('Đ', 'D')
    name_str = unicodedata.normalize('NFKD', name_str).encode('ascii', 'ignore').decode('utf-8')
    clean_filename = re.sub(r'[<>:"/\\|?*\x00-\x1F]', '_', name_str).rstrip('. ').strip() or 'export.xlsx'
    if not clean_filename.lower().endswith('.xlsx'):
        clean_filename += '.xlsx'
        
    styles_data = data.get('styles', {})
    col_widths_data = data.get('colWidths', {})
    merges_data = data.get('merges', {})
        
    wb = openpyxl.Workbook()
    wb.remove(wb.active)  # Remove default sheet
    
    thin_side = Side(border_style="thin", color="D1D5DB")
    border_all = Border(left=thin_side, right=thin_side, top=thin_side, bottom=thin_side)
    
    for sheet_name, sheet_data in data['sheets'].items():
        safe_title = "".join(ch for ch in sheet_name if ch not in r'\/*?:[]')[:31] or "Sheet"
        ws = wb.create_sheet(title=safe_title)
        sheet_styles = styles_data.get(sheet_name, {})
        
        for r_idx, row in enumerate(sheet_data):
            for c_idx, val in enumerate(row):
                cleaned_val = val
                if isinstance(val, str):
                    if val.startswith('='):
                        cleaned_val = val
                    elif val == "":
                        cleaned_val = None
                    else:
                        try:
                            if '.' in val:
                                cleaned_val = float(val)
                            else:
                                cleaned_val = int(val)
                        except ValueError:
                            cleaned_val = val
                
                cell = ws.cell(row=r_idx + 1, column=c_idx + 1, value=cleaned_val)
                
                cell_key = f"{r_idx},{c_idx}"
                if cell_key in sheet_styles:
                    c_style = sheet_styles[cell_key]
                    font_args = {}
                    if c_style.get('bold'): font_args['bold'] = True
                    if c_style.get('italic'): font_args['italic'] = True
                    if c_style.get('underline'): font_args['underline'] = 'single'
                    if c_style.get('fontSize'): font_args['size'] = int(c_style['fontSize'])
                    if c_style.get('fontFamily'): font_args['name'] = c_style['fontFamily']
                    if c_style.get('color'):
                        hex_col = c_style['color'].replace('#', '')
                        if len(hex_col) == 6:
                            font_args['color'] = hex_col
                    if font_args:
                        cell.font = Font(**font_args)
                        
                    if c_style.get('bg'):
                        hex_bg = c_style['bg'].replace('#', '')
                        if len(hex_bg) == 6:
                            cell.fill = PatternFill(start_color=hex_bg, end_color=hex_bg, fill_type='solid')
                            
                    align_args = {}
                    if c_style.get('align'): align_args['horizontal'] = c_style['align']
                    if c_style.get('valign'): align_args['vertical'] = c_style['valign']
                    if c_style.get('wrap'): align_args['wrap_text'] = True
                    if align_args:
                        cell.alignment = Alignment(**align_args)
                        
                    if c_style.get('border'):
                        cell.border = border_all

        if sheet_name in col_widths_data:
            for col_letter, width in col_widths_data[sheet_name].items():
                try:
                    ws.column_dimensions[col_letter].width = float(width)
                except Exception:
                    pass
                
        if sheet_name in merges_data:
            for merge_range in merges_data[sheet_name]:
                try:
                    if isinstance(merge_range, str):
                        ws.merge_cells(merge_range)
                    elif isinstance(merge_range, dict) and 's' in merge_range and 'e' in merge_range:
                        sr = merge_range['s']['r'] + 1
                        sc = merge_range['s']['c'] + 1
                        er = merge_range['e']['r'] + 1
                        ec = merge_range['e']['c'] + 1
                        ws.merge_cells(start_row=sr, start_column=sc, end_row=er, end_column=ec)
                    elif isinstance(merge_range, dict) and 'startRow' in merge_range:
                        sr = merge_range['startRow'] + 1
                        sc = merge_range['startCol'] + 1
                        er = merge_range['endRow'] + 1
                        ec = merge_range['endCol'] + 1
                        ws.merge_cells(start_row=sr, start_column=sc, end_row=er, end_column=ec)
                except Exception as me:
                    print(f"[Warn] merge_cells failed for {merge_range}: {me}")

        # Embed images if provided
        images_data = data.get('images', {})
        if sheet_name in images_data:
            from openpyxl.drawing.image import Image as OpenpyxlImage
            import base64
            for img_info in images_data[sheet_name]:
                try:
                    src = img_info.get('src', '')
                    if not src:
                        continue
                    if ',' in src:
                        src = src.split(',', 1)[1]
                    raw_bytes = base64.b64decode(src)
                    img_stream = io.BytesIO(raw_bytes)
                    op_img = OpenpyxlImage(img_stream)
                    if img_info.get('width'):
                        op_img.width = int(img_info['width'])
                    if img_info.get('height'):
                        op_img.height = int(img_info['height'])
                    
                    c = int(img_info.get('col', 0))
                    r = int(img_info.get('row', 0))
                    anchor = f"{get_column_letter(c + 1)}{r + 1}"
                    ws.add_image(op_img, anchor)
                except Exception as img_err:
                    print(f"[Warn] Failed to embed image on {sheet_name}: {img_err}")
                    
    return wb, clean_filename

def save_workbook_to_downloads(wb, filename):
    try:
        downloads_dir = os.path.join(os.path.expanduser('~'), 'Downloads')
        if not os.path.exists(downloads_dir):
            downloads_dir = os.path.join(os.environ.get('USERPROFILE', 'C:\\Users\\Admin'), 'Downloads')
        os.makedirs(downloads_dir, exist_ok=True)
        dest_path = os.path.join(downloads_dir, filename)
        wb.save(dest_path)
        print(f"[Export] Saved directly to Downloads: {dest_path}")
        return dest_path
    except Exception as e:
        print(f"[Export] Failed to auto-save to Downloads: {e}")
        return None

@app.route('/download', methods=['POST'])
def download_file():
    data = None
    if request.is_json:
        data = request.get_json()
    else:
        import json
        raw = request.form.get('payload') or request.form.get('data')
        if raw:
            try:
                data = json.loads(raw)
            except Exception:
                data = None

    if not data or 'sheets' not in data:
        return jsonify({'error': 'Invalid request data'}), 400
    
    try:
        wb, filename = generate_openpyxl_workbook(data)
        saved_path = save_workbook_to_downloads(wb, filename)
        output = io.BytesIO()
        wb.save(output)
        output.seek(0)
        
        resp = send_file(
            output,
            mimetype='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet',
            as_attachment=True,
            download_name=filename
        )
        if saved_path:
            resp.headers['X-Saved-Path'] = saved_path
        return resp
    except Exception as e:
        return jsonify({'error': f'Failed to generate Excel file: {str(e)}'}), 500

@app.route('/prepare_export', methods=['POST'])
def prepare_export():
    import json, uuid
    data = None
    if request.is_json:
        data = request.get_json()
    else:
        raw = request.form.get('payload') or request.form.get('data')
        if raw:
            try:
                data = json.loads(raw)
            except Exception:
                data = None

    if not data or 'sheets' not in data:
        return jsonify({'error': 'Invalid request data'}), 400
        
    try:
        wb, filename = generate_openpyxl_workbook(data)
        saved_path = save_workbook_to_downloads(wb, filename)
        token = uuid.uuid4().hex[:12]
        filepath = os.path.join(EXPORT_DIR, f"{token}_{filename}")
        wb.save(filepath)
        
        return jsonify({
            'success': True,
            'download_url': f'/get_file/{token}/{filename}',
            'filename': filename,
            'saved_to': saved_path or filepath
        })
    except Exception as e:
        return jsonify({'error': f'Failed to generate Excel file: {str(e)}'}), 500

@app.route('/get_file/<token>/<filename>', methods=['GET'])
def get_exported_file(token, filename):
    import glob
    matches = glob.glob(os.path.join(EXPORT_DIR, f"{token}_*"))
    if not matches:
        return "File not found or expired", 404
    filepath = matches[0]
    return send_file(
        filepath,
        mimetype='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet',
        as_attachment=True,
        download_name=filename
    )

if __name__ == '__main__':
    cli_target = get_target_file_from_args()
    if cli_target:
        app.config['INITIAL_FILE_PATH'] = cli_target
        print(f"Auto-opening spreadsheet file: {cli_target}")

    def open_browser():
        webbrowser.open('http://127.0.0.1:5000')
    threading.Timer(1.0, open_browser).start()
    app.run(host='0.0.0.0', debug=False, port=5000)
