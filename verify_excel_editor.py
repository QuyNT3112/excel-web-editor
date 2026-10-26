import io
import json
import unittest
import openpyxl
from app import app

class TestExcelWebEditor(unittest.TestCase):
    def setUp(self):
        self.app = app.test_client()
        self.app.testing = True

    def test_index_page(self):
        """Test if the main editor page renders successfully"""
        response = self.app.get('/')
        self.assertEqual(response.status_code, 200)
        self.assertIn(b'Excel Web Editor', response.data)
        self.assertIn(b'formula-bar', response.data)
        self.assertIn(b'sheet-tabs-bar', response.data)

    def test_excel_upload_and_download_flow(self):
        """Test uploading a dynamically generated Excel file, then exporting/downloading it back"""
        
        # 1. Create a dummy Excel workbook in memory
        wb = openpyxl.Workbook()
        ws1 = wb.active
        ws1.title = "SalesData"
        ws1.cell(row=1, column=1, value="Month")
        ws1.cell(row=1, column=2, value="Revenue")
        ws1.cell(row=2, column=1, value="January")
        ws1.cell(row=2, column=2, value=1500)
        ws1.cell(row=3, column=1, value="February")
        ws1.cell(row=3, column=2, value=2200)
        ws1.cell(row=4, column=1, value="Total")
        ws1.cell(row=4, column=2, value="=SUM(B2:B3)")
        
        # Add another sheet
        ws2 = wb.create_sheet(title="Notes")
        ws2.cell(row=1, column=1, value="Draft note")
        
        excel_buffer = io.BytesIO()
        wb.save(excel_buffer)
        excel_buffer.seek(0)

        # 2. Test upload endpoint
        upload_data = {
            'file': (excel_buffer, 'test_sales.xlsx')
        }
        
        response = self.app.post('/upload', data=upload_data, content_type='multipart/form-data')
        self.assertEqual(response.status_code, 200)
        
        json_data = json.loads(response.data.decode('utf-8'))
        self.assertEqual(json_data['filename'], 'test_sales.xlsx')
        self.assertIn('SalesData', json_data['sheets'])
        self.assertIn('Notes', json_data['sheets'])
        
        # Verify uploaded cell data structure
        sales_sheet = json_data['sheets']['SalesData']
        self.assertEqual(sales_sheet[0][0], "Month")
        self.assertEqual(sales_sheet[0][1], "Revenue")
        self.assertEqual(sales_sheet[1][1], 1500)
        self.assertEqual(sales_sheet[3][1], "=SUM(B2:B3)")
        
        # Verify Notes sheet
        notes_sheet = json_data['sheets']['Notes']
        self.assertEqual(notes_sheet[0][0], "Draft note")

        # 3. Modify data simulate client edits
        # Update February revenue to 3000, and set March data in Row 4
        sales_sheet[2][1] = "3000"  # Row 3 (index 2)
        sales_sheet[3][0] = "March"  # Row 4 (index 3)
        sales_sheet[3][1] = "4000"
        # Set Total formula in Row 5 (index 4)
        sales_sheet[4][0] = "Total"
        sales_sheet[4][1] = "=SUM(B2:B4)"
        
        # 4. Test download/export endpoint
        export_payload = {
            'filename': 'updated_sales.xlsx',
            'sheets': json_data['sheets']
        }
        
        download_response = self.app.post('/download', json=export_payload, content_type='application/json')
        self.assertEqual(download_response.status_code, 200)
        self.assertEqual(download_response.mimetype, 'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')
        
        # 5. Load returned bytes as workbook to verify content
        download_buffer = io.BytesIO(download_response.data)
        out_wb = openpyxl.load_workbook(download_buffer, data_only=False)
        
        out_ws = out_wb['SalesData']
        
        # Verify content conversions and formula retention
        self.assertEqual(out_ws.cell(row=1, column=1).value, "Month")
        self.assertEqual(out_ws.cell(row=2, column=2).value, 1500) # kept as integer
        self.assertEqual(out_ws.cell(row=3, column=2).value, 3000) # parsed to integer
        self.assertEqual(out_ws.cell(row=4, column=2).value, 4000) # parsed to integer
        self.assertEqual(out_ws.cell(row=5, column=2).value, "=SUM(B2:B4)") # formula saved correctly
        
        out_notes = out_wb['Notes']
        self.assertEqual(out_notes.cell(row=1, column=1).value, "Draft note")

    def test_open_with_initial_file_and_save_in_place(self):
        """Test the Open with / CLI argument workflow and saving back in-place"""
        import os, tempfile, sys
        from app import get_target_file_from_args

        # 1. Create a dummy test file simulating an Excel file opened from disk
        temp_dir = tempfile.mkdtemp()
        test_file = os.path.join(temp_dir, "MyReport.xlsx")
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "Q1_Report"
        ws['A1'] = "Category"
        ws['B1'] = "Amount"
        ws['A2'] = "Hardware"
        ws['B2'] = 5000
        ws['A3'] = "Services"
        ws['B3'] = 2500
        ws['A4'] = "Total"
        ws['B4'] = "=SUM(B2:B3)"
        wb.save(test_file)

        # 2. Test command-line argument parser
        orig_argv = sys.argv
        sys.argv = ['ExcelWebPro.exe', test_file]
        detected = get_target_file_from_args()
        self.assertEqual(detected, os.path.abspath(test_file))
        sys.argv = orig_argv

        # 3. Test initial file check endpoint
        app.config['INITIAL_FILE_PATH'] = test_file
        res = self.app.get('/api/initial-file')
        self.assertEqual(res.status_code, 200)
        info = json.loads(res.data.decode('utf-8'))
        self.assertTrue(info['has_file'])
        self.assertEqual(info['filename'], 'MyReport.xlsx')

        # 4. Test initial file load endpoint
        load_res = self.app.get('/api/load-initial-file')
        self.assertEqual(load_res.status_code, 200)
        data = json.loads(load_res.data.decode('utf-8'))
        self.assertIn('Q1_Report', data['sheets'])
        q1_sheet = data['sheets']['Q1_Report']
        self.assertEqual(q1_sheet[0][0], "Category")
        self.assertEqual(q1_sheet[1][1], 5000)
        self.assertEqual(q1_sheet[3][1], "=SUM(B2:B3)")

        # 5. Modify cell value and test save_in_place
        q1_sheet[1][1] = 6500
        save_payload = {
            'filename': 'MyReport.xlsx',
            'filepath': test_file,
            'sheets': data['sheets'],
            'styles': data.get('styles', {})
        }
        save_res = self.app.post('/api/save_in_place', json=save_payload)
        self.assertEqual(save_res.status_code, 200)
        save_data = json.loads(save_res.data.decode('utf-8'))
        self.assertTrue(save_data['success'])

        # Verify disk file has the updated value
        saved_wb = openpyxl.load_workbook(test_file, data_only=False)
        self.assertEqual(saved_wb['Q1_Report']['B2'].value, 6500)
        self.assertEqual(saved_wb['Q1_Report']['B4'].value, "=SUM(B2:B3)")

        # Cleanup
        app.config['INITIAL_FILE_PATH'] = None
        import shutil
        shutil.rmtree(temp_dir)
        print("SUCCESS: All backend and Open With verification tests passed successfully!")

    def test_image_upload_and_persistence(self):
        """Test uploading images, embedding images in worksheets, and loading them back"""
        import os, tempfile, base64, io, shutil
        from PIL import Image as PILImage
        from app import generate_openpyxl_workbook, parse_spreadsheet_data

        # 1. Create a tiny test PNG base64
        pil_img = PILImage.new('RGB', (80, 60), color='#10b981')
        buf = io.BytesIO()
        pil_img.save(buf, format='PNG')
        buf.seek(0)
        img_b64 = base64.b64encode(buf.read()).decode('ascii')
        data_url = f"data:image/png;base64,{img_b64}"

        temp_dir = tempfile.mkdtemp()
        test_file = os.path.join(temp_dir, "ImageTest.xlsx")

        payload = {
            'filename': 'ImageTest.xlsx',
            'filepath': test_file,
            'sheets': {
                'Sheet1': [["Photo", "Price"], ["Item 1", 100]]
            },
            'images': {
                'Sheet1': [{
                    'id': 'img-test-1',
                    'src': data_url,
                    'name': 'emerald_box.png',
                    'col': 2,
                    'row': 1,
                    'width': 80,
                    'height': 60
                }]
            }
        }

        # 2. Test generate_openpyxl_workbook with images
        wb, clean_fn = generate_openpyxl_workbook(payload)
        ws = wb['Sheet1']
        self.assertEqual(len(ws._images), 1)
        wb.save(test_file)

        # 3. Test parse_spreadsheet_data extracts the image correctly
        parsed = parse_spreadsheet_data(test_file, 'ImageTest.xlsx')
        self.assertIn('images', parsed)
        self.assertIn('Sheet1', parsed['images'])
        extracted_images = parsed['images']['Sheet1']
        self.assertEqual(len(extracted_images), 1)
        self.assertTrue(extracted_images[0]['src'].startswith('data:image/png;base64,'))
        self.assertEqual(extracted_images[0]['col'], 2)
        self.assertEqual(extracted_images[0]['row'], 1)

        shutil.rmtree(temp_dir)
        print("SUCCESS: Image upload, embedding, and extraction tests passed successfully!")

if __name__ == '__main__':
    unittest.main()
